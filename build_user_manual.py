from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "StockFlow_User_Manual.docx"

# compact_reference_guide preset with a restrained StockFlow blue override.
BLUE = "2E74B5"
DARK_BLUE = "0B2545"
INK = "172B4D"
MUTED = "5B6B7A"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F7FA"
LINE = "C9D4E0"
GOLD = "9A6B00"
RED = "9B1C1C"
FONT = "Aptos"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LINE, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_run_font(run, size=None, color=None, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para(paragraph, before=0, after=6, line=1.25, alignment=None, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment
    if keep:
        fmt.keep_with_next = True


def add_border_bottom(paragraph, color=LINE, size="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    for kind, value in (("begin", None), ("instrText", instruction), ("separate", None), ("text", "1"), ("end", None)):
        tag = "w:fldChar" if kind in ("begin", "separate", "end") else ("w:t" if kind == "text" else f"w:{kind}")
        node = OxmlElement(tag)
        if kind in ("begin", "separate", "end"):
            node.set(qn("w:fldCharType"), kind)
        else:
            node.text = value
        run._r.append(node)
    set_run_font(run, size=9, color=MUTED)


def add_para(doc, text="", style=None, size=11, color=INK, bold=False, italic=False, before=0, after=6, line=1.25, alignment=None):
    paragraph = doc.add_paragraph(style=style)
    set_para(paragraph, before, after, line, alignment)
    if text:
        set_run_font(paragraph.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items, level=0):
    style_name = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        paragraph = doc.add_paragraph(style=style_name)
        set_para(paragraph, 0, 4, 1.25)
        set_run_font(paragraph.add_run(item), size=11, color=INK)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        set_para(paragraph, 0, 4, 1.25)
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.20)
        set_run_font(paragraph.add_run(f"{index}.  {item}"), size=11, color=INK)


def add_callout(doc, label, text, fill=PALE_BLUE, label_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=LINE, size="8")
    p = cell.paragraphs[0]
    set_para(p, 1, 1, 1.2)
    set_run_font(p.add_run(f"{label}: "), size=10.5, color=label_color, bold=True)
    set_run_font(p.add_run(text), size=10.5, color=INK)
    add_para(doc, "", after=1)


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell, color=LINE, size="8")
        p = cell.paragraphs[0]
        set_para(p, 0, 0, 1.1)
        set_run_font(p.add_run(text), size=font_size, color=DARK_BLUE, bold=True)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            set_cell_border(cell)
            p = cell.paragraphs[0]
            set_para(p, 0, 0, 1.15)
            set_run_font(p.add_run(str(text)), size=font_size, color=INK)
    add_para(doc, "", after=1)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    for level, size, color, before, after in [(1, 16, BLUE, 18, 10), (2, 13, BLUE, 14, 7), (3, 12, "1F4D78", 10, 5)]:
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
    for style_name, left, hanging in [("List Bullet", 540, 270), ("List Bullet 2", 900, 270), ("List Number", 540, 270)]:
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Pt(left / 20)
        style.paragraph_format.first_line_indent = Pt(-hanging / 20)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.42)
        section.footer_distance = Inches(0.42)


def add_header_footer(doc):
    for section in doc.sections:
        p = section.header.paragraphs[0]
        p.text = ""
        set_para(p, 0, 2, 1.0)
        set_run_font(p.add_run("STOCKFLOW"), size=8.5, color=BLUE, bold=True)
        set_run_font(p.add_run("  |  User Manual"), size=8.5, color=MUTED)
        add_border_bottom(p)
        fp = section.footer.paragraphs[0]
        fp.text = ""
        set_para(fp, 2, 0, 1.0, WD_ALIGN_PARAGRAPH.RIGHT)
        set_run_font(fp.add_run("Internal reference  |  StockFlow Workspace   "), size=8.5, color=MUTED)
        add_field(fp, "PAGE")


def page_break(doc):
    doc.add_page_break()


def cover(doc):
    add_para(doc, "STOCKFLOW", size=11, color=BLUE, bold=True, before=18, after=18, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph(style="Title")
    set_para(p, 90, 8, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("User Manual")
    p = doc.add_paragraph()
    set_para(p, 0, 18, 1.15, WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(p.add_run("Inventory, Procurement & Logistics Workspace"), size=15, color=MUTED)
    p = doc.add_paragraph()
    set_para(p, 0, 22, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(p.add_run("—  OPERATIONS GUIDE  —"), size=10, color=BLUE, bold=True)
    add_para(doc, "A practical guide for administrators, procurement teams, storekeepers, accountants, department managers, and supplier accounts.", size=13, color=INK, after=24, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2400, 6960])
    metadata = [("Document version", "1.0"), ("Project snapshot", "18 August 2026"), ("Primary audience", "System users and operational administrators"), ("Related admin guide", "Cpanel.MD - production hosting reference")]
    for row, (label, value) in zip(table.rows, metadata):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_border(cell, color=LINE)
        set_run_font(row.cells[0].paragraphs[0].add_run(label), size=9.5, color=DARK_BLUE, bold=True)
        set_run_font(row.cells[1].paragraphs[0].add_run(value), size=9.5, color=INK)
    add_callout(doc, "How to use this manual", "Start with Sections 3 and 4 for access and navigation, then use the workflow chapter for your role. Administrators should also read Sections 7, 22, and 23.")
    add_para(doc, "Prepared for the StockFlow application", size=10, color=MUTED, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)


def contents(doc):
    add_heading(doc, "Contents", 1)
    add_para(doc, "This manual is organized by user task. Headings are native Word headings so the Navigation Pane can be used to jump between sections.", color=MUTED, after=12)
    rows = [
        ("1", "About StockFlow", "Purpose, scope, and operating model"),
        ("2", "System at a glance", "Modules and common records"),
        ("3", "Access, authentication, and security", "Login, lock screen, session behavior, and safe use"),
        ("4", "Navigation and common controls", "Sidebar, search, tables, filters, previews, and downloads"),
        ("5", "Dashboard", "Role-aware metrics, live operations, and quick actions"),
        ("6", "Organization structure", "Companies, branches, departments, and locations"),
        ("7", "Users, roles, and permissions", "Access administration and least privilege"),
        ("8", "Suppliers and supplier accounts", "Supplier records, contacts, terms, mappings, and portals"),
        ("9", "Supplier partnership documents", "Private supplier document workspaces"),
        ("10", "Procurement and purchase orders", "Create, approve, issue, receive, and pay"),
        ("11", "Purchase-order documents and PDF viewing", "Attachments, previews, downloads, and signatures"),
        ("12", "Petty cash", "Voucher issuance, PDFs, signatures, and controls"),
        ("13", "Inventory and stores", "Products, categories, units, stock, issues, and transfers"),
        ("14", "Warehouses and spare-part locator", "Location balances and workshop lookup"),
        ("15", "Department workspace", "Department orders and allocation"),
        ("16", "Calendar and announcements", "Role-relevant events and internal communication"),
        ("17", "Logistics", "Shipment visibility and delivery follow-up"),
        ("18", "Sales, POS, and customers", "Point of sale, sales history, and customer records"),
        ("19", "Reports and data exchange", "Operational reporting, import, and export"),
        ("20", "Media Library", "Shared organization media files"),
        ("21", "Profile, digital signature, and preferences", "Profile, signature, password, and theme"),
        ("22", "Administration, backups, and settings", "System maintenance and recovery"),
        ("23", "Troubleshooting and support", "Common problems and recovery steps"),
        ("24", "Hosting and release checklist", "cPanel deployment summary"),
        ("25", "Glossary and operational checklists", "Definitions and quick reference"),
    ]
    add_table(doc, ["Section", "Topic", "What it covers"], rows, [720, 3000, 5640], font_size=9)
    add_callout(doc, "Recommended reading paths", "Procurement staff: 3, 4, 8, 10, 11, and 19. Storekeepers: 4, 5, 13, and 14. Accountants: 3, 10, 12, and 19. Supplier users: 3, 8, 9, and 11. Administrators: all sections, especially 6, 7, 22, and 24.")
    page_break(doc)


def section_about(doc):
    add_heading(doc, "1. About StockFlow", 1)
    add_para(doc, "StockFlow is an organization workspace for inventory, procurement, logistics, supplier relationships, petty cash, sales, reporting, and operational administration. It combines a React web interface with a Laravel API and a role-aware permission model.")
    add_heading(doc, "What this manual covers", 2)
    add_bullets(doc, ["How users sign in, unlock a locked screen, navigate the workspace, and work safely with business records.", "How administrators configure organizations, users, roles, permissions, settings, backups, and hosting.", "How operational teams create and follow purchase orders, record receiving, verify payments, manage suppliers, and issue petty cash.", "How store and warehouse teams maintain products, locations, stock movements, and workshop spare-part information.", "How users preview, print, download, upload, and organize business documents."])
    add_heading(doc, "Operating principles", 2)
    add_table(doc, ["Principle", "What it means in practice"], [("One source of truth", "Use the relevant record page instead of maintaining parallel spreadsheets for active transactions."), ("Role-aware access", "Users see actions and modules based on their role and permissions. Do not share credentials to bypass an access limit."), ("Traceability", "Keep purchase-order attachments, comments, payment evidence, supplier documents, and stock movements with the record they explain."), ("Controlled change", "Approvals, receiving, payment verification, and destructive actions are separate steps to protect the audit trail.")], [2600, 6760])


def section_overview(doc):
    add_heading(doc, "2. System at a glance", 1)
    add_para(doc, "The system is organized into functional areas. The sidebar displays only modules available to the signed-in user, while the dashboard adapts its metrics and quick actions to the user role.")
    add_table(doc, ["Area", "Main records", "Typical users"], [("Overview", "Dashboard, calendar, announcements, notifications", "All users"), ("Administration", "Users, roles, user access, organization structure, settings, backups", "Super Admin and authorized administrators"), ("Inventory & Stores", "Products, categories, units, suppliers, customers, warehouses, stores", "Storekeepers, procurement, managers"), ("Procurement & Logistics", "Purchase orders, receiving, payments, petty cash, logistics", "Procurement, accountants, storekeepers"), ("Sales & Finance", "POS, sales history, reports", "Sales and finance users"), ("System & Content", "Media Library, supplier partnership documents", "Authorized users and supplier accounts")], [2200, 4300, 2860])
    add_heading(doc, "Record relationships", 2)
    add_para(doc, "A typical procurement chain is: department need -> purchase order -> supplier -> receiving location -> goods receipt -> payment evidence -> stock movement -> report. Supplier accounts connect to their organization purchase orders and their private partnership documents. A user profile can supply a visual signature used in generated PDFs.")
    add_callout(doc, "Important", "The application has empty states for records that have not yet been created. An empty chart or list is not necessarily an error; it may mean the underlying products, stock movements, events, or transactions have not been entered yet.")


def section_access(doc):
    add_heading(doc, "3. Access, authentication, and security", 1)
    add_heading(doc, "Signing in", 2)
    add_numbered(doc, ["Open the approved StockFlow application URL.", "Enter the email address and password assigned to your account.", "Complete any email verification or invitation acceptance step if prompted.", "After successful sign-in, confirm that your name and role are correct before starting work."])
    add_heading(doc, "Lock screen behavior", 2)
    add_para(doc, "The lock screen protects an active session when the workspace is locked. Enter the account password to unlock the existing session. The application should return you to the page where you were working. If an unlock redirects to login, refresh once and report the time, page, and browser to the administrator.")
    add_heading(doc, "Session and account safety", 2)
    add_bullets(doc, ["Never share an account or leave a signed-in workspace unattended.", "Use the lock screen when stepping away and sign out at the end of a shared-device session.", "Use a unique password with at least 10 characters including letters and numbers, and change it immediately if exposed.", "Confirm the supplier, receiving location, amount, and supporting document before submitting or approving a transaction.", "Treat uploaded documents as business records. Do not upload unrelated personal files or secrets."])
    add_heading(doc, "What to do after access is denied", 2)
    add_numbered(doc, ["Check that you are signed in to the correct account and organization.", "Check the page and action: viewing, creating, approving, managing, and deleting may require different permissions.", "Do not repeatedly retry destructive actions. Capture the message and contact an administrator.", "An administrator should review the user's role, permission overrides, supplier assignment, and organization assignment."])


def section_navigation(doc):
    add_heading(doc, "4. Navigation and common controls", 1)
    add_heading(doc, "Sidebar and page layout", 2)
    add_para(doc, "Use the left navigation to move between modules. On smaller screens, open the navigation menu from the application shell. The top area normally contains the page title, description, search or filters, and primary actions. The footer identifies the StockFlow workspace.")
    add_heading(doc, "Common controls", 2)
    add_table(doc, ["Control", "Purpose", "Good practice"], [("Search", "Find a record by name, number, supplier, SKU, or document text where supported.", "Use a specific term and clear the field when switching tasks."), ("Status filter", "Limit records to draft, sent, confirmed, received, cancelled, or other workflow states.", "Check the filter before assuming a record is missing."), ("Eye / View", "Open a profile, record, or in-system document preview depending on the page.", "Use the document preview for reading; use the record page for editing and workflow actions."), ("Download", "Save a PDF or uploaded file to the local device.", "Use only approved storage locations for downloaded business records."), ("Print", "Open a printable version of a supported document.", "Check the page preview before printing."), ("Pagination", "Move through a long list and adjust the visible row count where available.", "Use filters before increasing page size."), ("Confirmation modal", "Require a deliberate choice for deletion or other irreversible actions.", "Read the warning and confirm the scope before proceeding.")], [1800, 4200, 3360])
    add_heading(doc, "Document preview", 2)
    add_para(doc, "The shared document viewer opens PDFs, images, and text-supported files inside the application. Print and download icons appear in the top-right of the viewer. Office and archive files may display a download-only message when the browser cannot render them directly.")
    add_callout(doc, "If a document looks small", "Use the current application build and refresh after a deployment. The preview uses a wide responsive workspace and page-width PDF zoom; browser zoom or an old cached JavaScript bundle can make the viewer appear narrow.")


def section_dashboard(doc):
    add_heading(doc, "5. Dashboard", 1)
    add_para(doc, "The dashboard is a role-aware operating view. It presents the metrics and workflow stages most relevant to the signed-in user's work instead of showing the same cards to every user.")
    add_heading(doc, "Using dashboard cards", 2)
    add_bullets(doc, ["Read the value and label together; a value may represent count, currency, quantity, or a workflow stage.", "Select a card or quick action to open the related module with the relevant filter when provided.", "Use operational lists for the transaction-level detail behind a summary number.", "When a panel says there is no stock balance, warehouse balance, or movement history yet, create the underlying products and stock movements first."])
    add_heading(doc, "Role examples", 2)
    add_table(doc, ["Role", "Typical dashboard emphasis"], [("Super Admin", "Broad operational overview plus administration and system controls."), ("Procurement Manager", "Requests, approvals, purchase orders, supplier follow-up, and delayed deliveries."), ("Procurement Officer", "Approved demand, orders to issue, supplier response, and expected delivery."), ("Store Keeper", "Out-of-stock and low-stock products, receipts, issues, and transfers."), ("Accountant", "Orders awaiting approval, payment verification, verified payments, and outstanding value."), ("Department Manager", "Department orders, allocation, and request progress."), ("Supplier", "Supplier portal purchase orders and partnership document access.")], [2200, 7160])


def section_org(doc):
    add_heading(doc, "6. Organization structure", 1)
    add_para(doc, "Organization structure defines the company, branches, departments, warehouses, and stores used by records and permissions. Keep names and codes stable after transactions begin, because they are used in reports and allocation decisions.")
    add_heading(doc, "Recommended setup order", 2)
    add_numbered(doc, ["Create or verify the company profile, currency, contact details, and address.", "Create the head office or branches and mark active locations clearly.", "Create departments with codes and branch assignment.", "Create warehouses and stores with their branch and active status.", "Assign users to departments and, where applicable, supplier accounts.", "Confirm that purchase orders, inventory, and dashboard filters use the intended locations."])
    add_heading(doc, "Companies and branches", 2)
    add_para(doc, "The organization page groups branches and departments under companies. Use the search field to find a company, branch, or department. Edit actions change master data; delete actions should be used only when the record has no required history or when the system presents a safe removal option.")
    add_heading(doc, "Departments", 2)
    add_para(doc, "Departments are used for allocation and department workspace requests. When a purchase order is allocated to a department, use the current department record rather than typing a near-duplicate legacy name.")
    add_heading(doc, "Locations", 2)
    add_para(doc, "Warehouses and stores are receiving and stock locations. Keep codes unique, identify the branch, and deactivate a location rather than creating duplicate replacements when historical stock or purchase orders depend on it.")


def section_roles(doc):
    add_heading(doc, "7. Users, roles, and permissions", 1)
    add_para(doc, "Permissions control access to modules and actions. Roles provide reusable permission bundles. The project supports standard roles and administrator-created custom roles, plus user permission overrides where enabled.")
    add_table(doc, ["Role", "Primary responsibilities", "Typical access"], [("Super Admin", "System ownership and recovery.", "All modules, users, roles, settings, backups, destructive administration."), ("Procurement Manager", "Approve demand and purchase orders; oversee supplier and logistics work.", "Procurement approvals, suppliers, receiving view, reports, locations, calendar."), ("Procurement Officer", "Create and issue purchase orders; coordinate suppliers.", "Purchase-order creation and issue, master data, suppliers, receiving, reports."), ("Store Keeper", "Receive, inspect, post, issue, transfer, and count stock.", "Inventory, locations, goods receipt actions, warehouse locator, workshop stock."), ("Department Manager", "Submit and follow department requests.", "Department workspace, inventory view, organization context."), ("Accountant", "Verify payments and monitor financial records.", "Purchase orders, payment verification, reports, announcements."), ("Supplier", "Respond to supplier-side workflow and maintain partnership files.", "Supplier portal, linked orders, supplier document upload.")], [1900, 3360, 4100])
    add_heading(doc, "Managing access", 2)
    add_numbered(doc, ["Open Roles & Permissions to review the available permission catalogue and role bundles.", "Create or edit a custom role only when an existing standard role does not match the job.", "Assign the role to the user and verify the organization, department, or supplier assignment.", "Ask the user to sign out and back in, or refresh their session, after a material permission change.", "Review access periodically and remove unused accounts or permissions promptly."])
    add_callout(doc, "Least privilege", "Grant the smallest permission set that allows a user to perform their job. Viewing a purchase order does not automatically mean the user should create, approve, pay, or delete it.", fill="FFF8E7", label_color=GOLD)


def section_suppliers(doc):
    add_heading(doc, "8. Suppliers and supplier accounts", 1)
    add_para(doc, "The supplier command center keeps legal details, contacts, commercial terms, product mappings, documents, and transaction history in one account.")
    add_heading(doc, "Create or update a supplier", 2)
    add_numbered(doc, ["Open Suppliers and choose Add Supplier, or open an existing supplier and choose Edit.", "Enter the legal name, address, city, country, email, phone, tax identifiers, currency, payment terms, and relationship notes.", "Save and confirm the supplier appears in the supplier list with the intended active status.", "Open the supplier profile to add contacts, commercial terms, product catalog mappings, and partnership documents."])
    add_heading(doc, "Supplier profile areas", 2)
    add_table(doc, ["Area", "Use it for"], [("General Information", "Legal identity, TIN / tax IDs, currency, contact channels, and internal notes."), ("Contact Roster", "Sales, support, accounting, and primary supplier contacts."), ("Commercial Terms", "Payment terms, minimum order quantity, preferred shipping, pricing tiers, and volume discounts."), ("Product Catalog Mapping", "Supplier SKU and part number mapped to internal products and unit conversion."), ("Document Vault", "Existing supplier records, expiry monitoring, preview, and download."), ("Transaction History", "Purchase orders, returns, invoices, status, dates, and amounts."), ("Partnership Documents", "Dedicated private workspace for contracts, agreements, certificates, and related files.")], [3000, 6360])
    add_heading(doc, "Deleting a supplier", 2)
    add_para(doc, "Use the delete icon in the supplier list only after reviewing the warning. If purchase orders exist, the system should ask whether to remove the supplier with its PO history or keep the purchase orders and related records. Choose the option that preserves the required audit history.")


def section_supplier_docs(doc):
    add_heading(doc, "9. Supplier partnership documents", 1)
    add_para(doc, "Each supplier account has a private partnership document workspace. Internal authorized users can manage files from the supplier profile. Supplier accounts can access their own workspace from the Supplier Portal when the supplier assignment is present.")
    add_heading(doc, "Upload a partnership document", 2)
    add_numbered(doc, ["Open the supplier profile and choose Partnership documents, or use the same action in the Supplier Portal.", "Choose Upload document.", "Select a type: Contract / agreement, Compliance certificate, Insurance document, or Other partnership document.", "Enter a clear title, expiry date if applicable, and notes that explain the file's purpose.", "Select a PDF, image, Office, CSV, or text file up to 25 MB and submit the upload.", "Open the new file with the eye icon to confirm it can be read, then keep the original record name meaningful."])
    add_heading(doc, "Document controls", 2)
    add_bullets(doc, ["Search by title, original filename, type, or notes.", "Switch between grid and list views to scan many files or inspect details.", "Use the preview viewer for supported files; print and download controls are in the viewer header.", "Use the expiry badge to identify valid, expiring, or expired files.", "Delete only obsolete or incorrect files and confirm the deletion warning."])
    add_callout(doc, "Privacy", "Supplier documents are scoped to the supplier account and organization. Supplier-linked users are restricted to their assigned supplier. Do not place a document for one supplier in another supplier's workspace.")


def section_procurement(doc):
    add_heading(doc, "10. Procurement and purchase orders", 1)
    add_para(doc, "Purchase orders move through a controlled workflow. The exact available actions depend on status and permission.")
    add_heading(doc, "Purchase-order lifecycle", 2)
    add_table(doc, ["Stage", "Meaning", "Typical next action"], [("Draft", "Order is being prepared or awaits approval.", "Edit, approve, send, or cancel according to permission."), ("Sent", "Order has been issued or is awaiting confirmation.", "Approve, confirm, cancel, or follow up with supplier."), ("Confirmed", "Supplier has confirmed the order.", "Receive goods when delivered; record payment when authorized."), ("Partial", "Some but not all ordered quantity has been received.", "Record another receipt or complete the receipt."), ("Received", "Expected goods have been received.", "Complete payment and reporting actions."), ("Cancelled", "Order is closed without continuing the workflow.", "Retain the record for history and reporting.")], [1500, 4100, 3760])
    add_heading(doc, "Create a purchase order", 2)
    add_numbered(doc, ["Open Purchases and choose Create purchase order.", "Select the supplier, receiving location, expected delivery date, purchase reason, and department allocation.", "Add each product, ordered quantity, unit cost, and tax rate. Confirm line totals and the overall total.", "Add collection or ownership information and notes when required.", "Save the order as a draft and review the supplier, location, dates, items, tax, and totals.", "Use the approval or issue action allowed by your role. The system may require a manager approval before an approved LPO PDF is available."])
    add_heading(doc, "Receive goods", 2)
    add_numbered(doc, ["Open a confirmed or partially received purchase order.", "Compare the delivery to the order line by line.", "Enter the received quantity, attach a receipt or delivery evidence file, and submit the receipt.", "Check the updated received quantity and stock movement before closing the delivery activity."])
    add_heading(doc, "Verify payment", 2)
    add_para(doc, "An authorized accountant can record a payment for an approved, non-draft order. Select the outstanding purchase order, enter the amount, attach proof of payment, and verify the result. Never mark a payment as verified without matching evidence.")


def section_po_docs(doc):
    add_heading(doc, "11. Purchase-order documents and PDF viewing", 1)
    add_para(doc, "Purchase orders support both generated documents and uploaded record evidence. The same preview experience is available from the procurement list, purchase-order detail page, supplier portal, and attachment sections.")
    add_heading(doc, "Generated LPO PDF", 2)
    add_bullets(doc, ["The eye icon previews an approved LPO PDF in the system.", "The download icon saves the approved LPO PDF.", "Draft or unapproved orders intentionally show an approval requirement instead of an approved document action.", "The PDF may include a visual signature from the authorized user's profile and the company information used by the generator."])
    add_heading(doc, "Attachments", 2)
    add_table(doc, ["Attachment group", "Examples", "Action"], [("Receipts", "Goods receipt notes, delivery acknowledgements, receiving evidence.", "Upload, preview, download, remove when authorized."), ("Invoices", "Supplier invoices and financial support.", "Upload, preview, download, remove when authorized."), ("Proofs of payment", "Payment verification evidence.", "Preview and download; deletion is restricted by workflow.")], [1900, 4860, 2600])
    add_heading(doc, "Preview troubleshooting", 2)
    add_bullets(doc, ["If the file is not supported by the browser, use the download icon and open it with the appropriate desktop application.", "If the viewer is narrow, refresh after deployment and check browser zoom. The viewer uses a wide responsive layout and page-width zoom for PDFs.", "If a download or preview fails, verify that the account still has permission and that the API and storage link are available."])


def section_petty(doc):
    add_heading(doc, "12. Petty cash", 1)
    add_para(doc, "Petty cash vouchers record small cash issues with a purpose, beneficiary, issuer, amount, currency, notes, and a generated control-copy PDF.")
    add_heading(doc, "Issue a voucher", 2)
    add_numbered(doc, ["Open Petty Cash and choose Issue Petty Cash.", "Enter the date, collector / beneficiary, reason for voucher, amount, and optional notes.", "Confirm that the amount is positive and the purpose is specific enough for later review.", "Submit the voucher. The system creates the record and prepares the PDF.", "Use the eye icon to preview the PDF or the download icon to save it. Print from the viewer when a paper control copy is required."])
    add_heading(doc, "Petty cash controls", 2)
    add_bullets(doc, ["Do not use petty cash as a substitute for a purchase order when a procurement workflow is required.", "Use a clear reason instead of a generic description such as 'miscellaneous'.", "Keep the voucher number on any supporting receipt or reimbursement record.", "Check the issued-by and collector fields before distributing the document."])


def section_inventory(doc):
    add_heading(doc, "13. Inventory and stores", 1)
    add_para(doc, "Inventory records the product catalogue and stock movements that support receiving, issuing, transferring, counting, and replenishment decisions.")
    add_heading(doc, "Master data", 2)
    add_table(doc, ["Record", "Purpose"], [("Products", "Name, SKU, category, unit, reorder level, pricing, and active status."), ("Categories / subcategories", "Group products for searching, reporting, and stock review."), ("Brands", "Identify product manufacturer or brand where used."), ("Units", "Define how quantities are measured."), ("Suppliers", "Connect product sourcing and purchase orders to a supplier account."), ("Locations", "Identify where stock is received, stored, issued, or transferred.")], [2600, 6760])
    add_heading(doc, "Stock operations", 2)
    add_table(doc, ["Operation", "Use when", "Control"], [("Adjustment", "Correct an approved stock discrepancy or opening balance.", "Record a reason and use only the adjustment permission."), ("Issue", "Goods leave a warehouse or store for a department, workshop, or user.", "Confirm item, quantity, source, and collector."), ("Transfer", "Move goods between authorized locations.", "Confirm dispatch and receiving locations and the transfer quantity."), ("Count", "Compare physical stock to system quantity.", "Perform counts routinely and investigate differences."), ("Replenishment request", "A location needs procurement support for stock.", "Use when reorder or out-of-stock signals require action.")], [1800, 4420, 3140])
    add_heading(doc, "Good stock practice", 2)
    add_bullets(doc, ["Use one product record per SKU and avoid duplicate names for the same item.", "Post receipts before issuing goods from the receiving location.", "Use the dashboard's stock health and warehouse comparison panels as a prompt to investigate, not as a replacement for the ledger.", "Keep stock movement descriptions and reference numbers clear enough for an auditor to follow."])


def section_warehouses(doc):
    add_heading(doc, "14. Warehouses and spare-part locator", 1)
    add_heading(doc, "Warehouses and stores", 2)
    add_para(doc, "Use the storage locations module to review active warehouses and stores, their branches, and their relationship to stock. Open a location to view stock balances and location statistics where available.")
    add_heading(doc, "Spare-part locator", 2)
    add_para(doc, "The warehouse locator helps workshop users find where a spare part is stored. Search by product or SKU, confirm the location and quantity, and update the locator only when authorized. Locator updates do not replace stock adjustments or transfer records.")
    add_callout(doc, "Avoid double counting", "A locator is a reference for finding a part. The stock ledger remains the authoritative record for quantity. Use transfers and adjustments to change stock, not the locator alone.")


def section_department(doc):
    add_heading(doc, "15. Department workspace", 1)
    add_para(doc, "Department Workspace gives departments a structured way to submit and follow requests that may become procurement activity. Allocation fields on purchase orders connect the request to the department receiving the value.")
    add_heading(doc, "Submit a department order", 2)
    add_numbered(doc, ["Open Department Orders and review existing requests before creating a duplicate.", "Enter the requested products or services, purpose, quantities, and required date.", "Submit the order for the department's approval and procurement follow-up path.", "Monitor the request status and communicate clarifications through the available record context."])
    add_para(doc, "Department managers should use the workspace for planned demand and use the purchase order pages for the formal supplier-facing order. Keep the purpose and allocation consistent between both records.")


def section_calendar(doc):
    add_heading(doc, "16. Calendar and announcements", 1)
    add_heading(doc, "Calendar", 2)
    add_para(doc, "The calendar is designed around events relevant to the user's account, role, and available procurement or organizational data. It is not intended to display unrelated holiday, leave, or birthday categories unless those records are part of the configured system scope.")
    add_bullets(doc, ["Review upcoming purchase-order delivery dates, approvals, receiving activities, meetings, and other configured operational events.", "Use the month summary to understand the number of events visible to the current user.", "If no events appear, confirm that the underlying purchase orders, locations, calendar permissions, and date fields exist."])
    add_heading(doc, "Announcements", 2)
    add_para(doc, "Announcements provide internal communication. High-priority announcements may create in-app alerts for active system users. Read the title, message, priority, and any audience or timing information before acting.")


def section_logistics(doc):
    add_heading(doc, "17. Logistics", 1)
    add_para(doc, "Logistics provides shipment and delivery visibility for operational follow-up. Use it alongside expected delivery dates on purchase orders and receiving records.")
    add_heading(doc, "Delivery follow-up", 2)
    add_numbered(doc, ["Open Logistics and filter to active or delayed shipments.", "Compare the expected date, supplier, destination, and shipment status against the purchase order.", "Record or communicate a follow-up through the available workflow and keep evidence with the related record.", "When goods arrive, complete receiving in the purchase-order workflow so stock and payment processes can continue."])


def section_sales(doc):
    add_heading(doc, "18. Sales, POS, and customers", 1)
    add_heading(doc, "Point of Sale", 2)
    add_para(doc, "Authorized sales users use POS to create sales against active products and available stock. Confirm the product, quantity, selling price, discount, payment method, and customer details before completing a sale.")
    add_heading(doc, "Sales history", 2)
    add_para(doc, "Use Sales History to search completed sales, inspect amounts and payment information, and support daily reconciliation. Do not void or correct a sale without the required permission and a documented reason.")
    add_heading(doc, "Customers", 2)
    add_para(doc, "Customer records store contact and relationship information used by sales operations. Keep names and contact details accurate, avoid duplicates, and protect customer data when exporting or downloading reports.")


def section_reports(doc):
    add_heading(doc, "19. Reports and data exchange", 1)
    add_heading(doc, "Reports", 2)
    add_table(doc, ["Report family", "Questions it helps answer"], [("Procurement / purchases", "How many purchase orders exist, what has been spent, what is outstanding, and how spend is distributed."), ("Stock / inventory", "What is in stock, what is low or out of stock, and how inventory is valued."), ("Sales", "What was sold, revenue, transaction count, discounts, and product performance."), ("Payments and returns", "What has been paid, what is awaiting verification, and what returns have been recorded.")], [3000, 6360])
    add_para(doc, "Use filters and date ranges carefully. A report is only as accurate as the underlying master data, stock movements, purchase orders, sales records, and payment evidence.")
    add_heading(doc, "Data export and import", 2)
    add_bullets(doc, ["Export only the data necessary for the business task and store the file securely.", "Review CSV or spreadsheet columns before importing. Keep identifiers, dates, quantities, and currency values consistent.", "Use import authorization when the system requests it and verify the result after the transaction completes.", "Never use import to bypass workflow approvals or overwrite history without a documented correction process."])


def section_media(doc):
    add_heading(doc, "20. Media Library", 1)
    add_para(doc, "Media Library is the shared organization file area for images and general media files. Supplier partnership documents should normally remain in the supplier-specific workspace instead.")
    add_heading(doc, "Upload media", 2)
    add_numbered(doc, ["Open Media Library and choose Upload Media.", "Select one or more files. Supported categories include images, PDF, Office files, CSV, text, and ZIP files within the displayed limits.", "Choose a folder when folders are configured and submit the upload.", "Confirm the file appears in the grid or list and use the eye icon to open it in the shared viewer."])
    add_heading(doc, "Folders and cleanup", 2)
    add_bullets(doc, ["Use folders for stable organizational groupings such as branding, templates, operations, and reference material.", "Use search before uploading a duplicate file.", "Delete only files that are obsolete or uploaded in error; deletion is irreversible in the user interface.", "Use supplier-specific document workspaces for files that should not be shared across the organization media area."])


def section_profile(doc):
    add_heading(doc, "21. Profile, digital signature, and preferences", 1)
    add_heading(doc, "Profile information", 2)
    add_para(doc, "Open your profile to update your name, sign-in email, phone number, avatar, digital signature, and password. Changing the sign-in email requires the current password.")
    add_heading(doc, "Upload a digital signature", 2)
    add_numbered(doc, ["Open Profile and locate the Digital Signature field.", "Choose a PNG file. A transparent-background PNG is recommended so the signature sits naturally on a document.", "Confirm the preview looks correct and select Save.", "Use the signature in an approved generated PDF to confirm the visual signature appears in the designated signature field."])
    add_callout(doc, "Signature troubleshooting", "The system accepts PNG signature files up to 1 MB. If a valid PNG is rejected, hard-refresh the application so the current multipart upload code is loaded, then reselect the file and save again. Do not rename a non-PNG file to .png.", fill="FFF8E7", label_color=GOLD)
    add_heading(doc, "Theme and preferences", 2)
    add_para(doc, "Use the account and settings controls available to your role to change supported preferences such as theme, date display, currency configuration, notification behavior, and organization settings. Always test a setting change on a non-destructive page before applying a system-wide change.")


def section_admin(doc):
    heading = add_heading(doc, "22. Administration, backups, and settings", 1)
    heading.paragraph_format.page_break_before = True
    add_heading(doc, "System settings", 2)
    add_bullets(doc, ["Keep the company name, address, phone, email, currency, timezone, and date format accurate because they appear in documents and reports.", "Configure notification and email settings only after confirming the production mail service.", "Use the cache-clear action after a controlled configuration change when the application provides it."])
    add_heading(doc, "Backups", 2)
    add_numbered(doc, ["Open Backups / Settings as an authorized administrator.", "Set the frequency, retention period, storage destination, email copy option, and whether uploads are included.", "Run a controlled backup and record the filename and timestamp.", "Periodically verify that a backup can be located and that its contents are usable; a backup that has never been tested is not a recovery plan."])
    add_heading(doc, "Release cache sequence", 2)
    add_para(doc, "After a backend release, use the documented Laravel optimization sequence. After a frontend release, publish the new hashed dist assets and ensure index.html is not served from a stale cache. See Section 24 for the hosting checklist.")
    add_heading(doc, "Destructive administration", 2)
    add_callout(doc, "Before deletion", "Confirm the exact target, review linked purchase orders and history, make a backup when material records are involved, and use the system's warning modal. Do not delete broad folders or database collections from a production host.", fill="FDECEC", label_color=RED)


def section_troubleshoot(doc):
    add_heading(doc, "23. Troubleshooting and support", 1)
    add_table(doc, ["Symptom", "Likely cause", "First response"], [("Page will not display", "Stale frontend bundle, route error, API failure, or a rendering exception.", "Hard-refresh, check the browser console, confirm the API is reachable, and retry the direct route."), ("Unlock returns to login", "Expired session, invalid unlock response, or stale session state.", "Sign in again, preserve the return URL, and report the page and time if it repeats."), ("Supplier PO page is blank", "Missing PO record, API permission, or frontend data shape mismatch.", "Open the PO from the list, confirm the ID, check API response, and verify supplier order permission."), ("Document preview is narrow", "Old cached bundle, browser zoom, or embedded PDF viewer settings.", "Refresh, reset browser zoom, and reopen the file. The current viewer uses 98% responsive width."), ("Signature upload rejected", "Old JavaScript bundle, incorrect file type, or multipart request issue.", "Hard-refresh, reselect a real PNG under 1 MB, and save. Do not use a renamed JPG."), ("Cannot delete supplier", "Purchase orders or related records are linked.", "Use the supplier deletion choice that preserves the required PO history."), ("No dashboard values", "Underlying products, stock movements, warehouses, or transactions do not exist.", "Set up master data and record the relevant operations before expecting live values."), ("Upload fails", "File limit, storage permissions, missing storage link, or API error.", "Check file type/size, storage link, Laravel logs, PHP limits, and directory permissions.")], [2200, 3500, 3660], font_size=9)
    add_heading(doc, "Information to include in a support report", 2)
    add_bullets(doc, ["Your user name, role, and organization (never include your password or token).", "The exact page URL and action you attempted.", "The approximate time, browser, device, and whether a hard refresh changed the result.", "The complete error message and a screenshot with sensitive information hidden.", "The record number or filename involved, if applicable."])


def section_hosting(doc):
    add_heading(doc, "24. Hosting and release checklist", 1)
    add_para(doc, "The project includes Cpanel.MD with the complete hosting guide. This section summarizes the user-visible operational checks for a production deployment. The backend should be hosted as Laravel with backend/public as the web root; the frontend should serve the built dist files.")
    add_heading(doc, "Recommended production layout", 2)
    add_table(doc, ["Component", "Recommended location", "Key rule"], [("Frontend", "app.example.com -> uploaded dist/ contents", "Serve index.html with SPA rewrite and cache hashed assets."), ("API", "api.example.com -> backend/public/", "Never expose the Laravel project root, .env, storage/, or vendor/."), ("Database", "cPanel MySQL / MariaDB", "Use a dedicated database user with required privileges."), ("Storage", "Laravel storage with public link", "Ensure uploads, signatures, avatars, and documents can be read by the application.")], [1800, 3900, 3660])
    add_heading(doc, "Deployment sequence", 2)
    add_numbered(doc, ["Back up the database and retain the previous frontend and backend release directories.", "Build the frontend with the production API URL; run typecheck and build before upload.", "Install backend dependencies with Composer in production mode and set APP_DEBUG=false.", "Configure HTTPS, database credentials, frontend origin, secure cookies, upload limits, and CORS.", "Run migrations with --force, create the storage link, and apply storage/bootstrap/cache permissions.", "Clear and rebuild Laravel configuration, route, view, and event caches.", "Verify login, direct frontend routes, lock-screen unlock, uploads, document preview, PDF downloads, and API health.", "Configure queue and scheduler cron jobs if mail, notifications, or scheduled tasks are enabled."])
    add_heading(doc, "cPanel production checklist", 2)
    add_bullets(doc, ["PHP 8.2+, Composer 2, Node.js 20+, MySQL 8 or MariaDB 10.6+, and required PHP extensions are available.", "Both frontend and API domains use HTTPS and CORS allows only the production frontend origin.", "index.html is no-cache; hashed JS, CSS, and image assets use long-lived immutable caching.", "PHP upload_max_filesize and post_max_size support supplier documents and PDF evidence.", "Laravel logs, PHP errors, disk space, MySQL slow queries, queue logs, and cron output are monitored.", "A rollback copy and database backup exist before each release."])


def section_glossary(doc):
    add_heading(doc, "25. Glossary and operational checklists", 1)
    add_heading(doc, "Glossary", 2)
    add_table(doc, ["Term", "Meaning"], [("Approved LPO", "The generated local purchase-order PDF made available after the order reaches the required approval state."), ("Allocation", "The department or organizational destination assigned to a purchase need or order."), ("Goods receipt", "The recorded quantity of goods physically received against an order."), ("Supplier account", "A supplier-linked user account that can access the supplier portal and its private documents."), ("Supplier document", "A file stored against one supplier, such as a contract, compliance certificate, insurance record, or agreement."), ("Media asset", "A shared organization file stored in the Media Library."), ("Permission override", "A user-specific allow or deny that changes the effective role permission where enabled."), ("Stock adjustment", "A controlled change to stock quantity used to correct an approved discrepancy or opening balance.")], [2600, 6760])
    add_heading(doc, "Daily operations checklist", 2)
    add_bullets(doc, ["Review dashboard alerts, approvals, delayed deliveries, low stock, and payment verification tasks.", "Confirm new supplier and department data before creating dependent transactions.", "Review purchase-order status and expected delivery dates.", "Complete receiving and attach evidence when goods arrive.", "Verify payment evidence before recording or approving payments.", "Use the shared viewer to confirm generated PDFs and uploaded documents are readable."])
    add_heading(doc, "Monthly administration checklist", 2)
    add_bullets(doc, ["Review active users, role assignments, supplier-linked accounts, and unused permissions.", "Review supplier document expiry badges and request renewals before expiration.", "Review backup status, retention, storage space, and restore evidence.", "Review stock adjustments, transfers, issues, and outstanding purchase orders.", "Review application and hosting logs for repeated errors, failed uploads, CORS errors, and stale cache reports.", "Document releases, configuration changes, and rollback points."])
    add_callout(doc, "End of manual", "Keep this guide with the deployment notes and update it when roles, workflow states, upload limits, hosting layout, or major modules change.")


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    props = doc.core_properties
    props.title = "StockFlow User Manual"
    props.subject = "Inventory, procurement, logistics, supplier, and administration user guide"
    props.author = "StockFlow"
    props.keywords = "StockFlow, inventory, procurement, logistics, user manual"
    cover(doc)
    contents(doc)
    for section in [section_about, section_overview, section_access, section_navigation, section_dashboard, section_org, section_roles, section_suppliers, section_supplier_docs, section_procurement, section_po_docs, section_petty, section_inventory, section_warehouses, section_department, section_calendar, section_logistics, section_sales, section_reports, section_media, section_profile, section_admin, section_troubleshoot, section_hosting, section_glossary]:
        section(doc)
    add_header_footer(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
