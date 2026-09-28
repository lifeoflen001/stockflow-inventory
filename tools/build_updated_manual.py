from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\xampp\htdocs\Inventory")
OUTPUT = Path(r"C:\Users\Admin\Downloads\Procurement_Sys_User_Manual.docx")
LOGO = ROOT / "src" / "assets" / "logo3.png"

BLUE = "2457A6"
BLUE_DARK = "173B79"
BLUE_LIGHT = "EAF1FC"
INK = "17233D"
MUTED = "5D6B85"
LINE = "D9E2F0"
PALE = "F6F9FD"
WHITE = "FFFFFF"
AMBER = "A86600"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs[edge]
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn("w:" + key), str(edge_data[key]))


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def font_run(run, size=None, color=INK, bold=False, italic=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, before=0, after=6, line=1.08, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_text(doc, text, style="Normal", before=0, after=6, color=INK, bold=False, italic=False, size=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    style_paragraph(p, before, after, keep=style.startswith("Heading"))
    run = p.add_run(text)
    font_run(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading " + str(level))
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    font_run(run, size=16 if level == 1 else 11.5, color=BLUE_DARK if level == 1 else BLUE, bold=True)
    return p


def add_bullets(doc, items, color=INK):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, after=3)
        font_run(p.add_run(item), size=9.6, color=color)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        style_paragraph(p, after=4)
        font_run(p.add_run(item), size=9.6, color=INK)


def add_callout(doc, title, body, fill=BLUE_LIGHT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 6.9)
    set_cell_shading(cell, fill)
    set_cell_border(cell, left={"val": "single", "sz": 18, "color": accent}, top={"val": "single", "sz": 4, "color": LINE}, bottom={"val": "single", "sz": 4, "color": LINE}, right={"val": "single", "sz": 4, "color": LINE})
    set_cell_margins(cell, top=150, start=190, bottom=150, end=190)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    style_paragraph(p, after=3)
    font_run(p.add_run(title), size=10, color=accent, bold=True)
    p2 = cell.add_paragraph()
    style_paragraph(p2, after=0)
    font_run(p2.add_run(body), size=9.4, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths=None, header_fill=BLUE_DARK):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    if widths is None:
        widths = [6.9 / len(headers)] * len(headers)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, top={"val": "single", "sz": 6, "color": header_fill}, bottom={"val": "single", "sz": 6, "color": header_fill}, left={"val": "single", "sz": 4, "color": WHITE}, right={"val": "single", "sz": 4, "color": WHITE})
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        style_paragraph(p, after=0)
        font_run(p.add_run(header), size=8.8, color=WHITE, bold=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            set_cell_shading(cell, PALE if row_index % 2 == 0 else WHITE)
            set_cell_border(cell, top={"val": "single", "sz": 3, "color": LINE}, bottom={"val": "single", "sz": 3, "color": LINE}, left={"val": "single", "sz": 3, "color": LINE}, right={"val": "single", "sz": 3, "color": LINE})
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, after=0)
            font_run(p.add_run(str(value)), size=8.8, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    font_run(run, size=8, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Heading 1", 16, BLUE_DARK), ("Heading 2", 11.5, BLUE), ("Heading 3", 10, BLUE)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(9.6)


def configure_page(doc):
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.header_distance = Inches(0.3)
        section.footer_distance = Inches(0.3)
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_paragraph(hp, after=0)
        font_run(hp.add_run("PROCUREMENT SYS  /  USER MANUAL"), size=7.5, color=MUTED, bold=True)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(fp, after=0)
        font_run(fp.add_run("Procurement Sys  •  Internal operations guide  •  "), size=8, color=MUTED)
        add_page_field(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=12)
    if LOGO.exists():
        logo_run = p.add_run()
        logo_run.add_picture(str(LOGO), width=Inches(2.25))
        doc_pr = logo_run._r.xpath(".//wp:docPr")[0]
        doc_pr.set("descr", "Procurement Sys company logo")
    add_text(doc, "PROCUREMENT SYS", before=10, after=2, color=BLUE_DARK, bold=True, size=26, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "User Manual", after=4, color=BLUE, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "A short, practical guide for daily inventory, procurement, supplier, logistics, petty cash, and administration work.", after=18, color=MUTED, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_callout(doc, "READ THIS FIRST", "Start with Quick start, then use the role path that matches your work. You only need the sections relevant to your account.", fill=BLUE_LIGHT, accent=BLUE)
    add_table(doc, ["Manual details", "Value"], [
        ("Release focus", "Current implemented workflows and controls"),
        ("Audience", "All users, supplier accounts, and administrators"),
        ("System name", "Procurement Sys"),
        ("Last updated", "19 August 2026"),
    ], widths=[1.75, 5.15], header_fill=BLUE)
    add_text(doc, "The application is role-aware. If a button or module is not visible, it is usually controlled by your account, organization, status, or permission.", before=12, after=0, color=MUTED, italic=True, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_contents(doc):
    add_heading(doc, "Contents", 1)
    add_text(doc, "Use the list below as a quick map. The manual is intentionally task-focused: follow the shortest path to the action you need.", after=10, color=MUTED)
    rows = [
        ("1", "Quick start", "Sign in, unlock, and choose your role path"),
        ("2", "System structure", "How the React, Laravel, database, files, and jobs fit together"),
        ("3", "Security and account safety", "Authentication, permissions, sessions, uploads, and audit controls"),
        ("4", "Navigation and dashboard", "Sidebar, search, empty states, notifications, and calendar"),
        ("5", "Organization and users", "Companies, branches, departments, locations, roles, and access"),
        ("6", "Suppliers and supplier portal", "Supplier records, products/prices, documents, and deletion choices"),
        ("7", "Procurement and LPOs", "Requests, purchase orders, receiving, payments, PDFs, and comments"),
        ("8", "Petty cash", "Vouchers, approvals, signatures, previews, and records"),
        ("9", "Inventory and workshop", "Products, stock movements, warehouse issues, and locator"),
        ("10", "Documents and previews", "Media Library, attachments, print, download, and supported files"),
        ("11", "Offline mode and synchronization", "Pending operations, conflicts, retries, and recovery"),
        ("12", "Policies and administration", "Policy pages, settings, notifications, backups, and deletion"),
        ("13", "Troubleshooting", "Common messages and first responses"),
    ]
    add_table(doc, ["#", "Section", "Use it for"], rows, widths=[0.35, 2.2, 4.35], header_fill=BLUE_DARK)
    add_callout(doc, "FASTEST ROUTE", "Most users need Sections 1, 4, and the workflow section for their role. Administrators should also read Sections 2, 3, 5, 11, and 12.", fill="FFF7E8", accent=AMBER)
    doc.add_page_break()


def section_quick_start(doc):
    add_heading(doc, "1. Quick start", 1)
    add_text(doc, "Procurement Sys opens with a role-aware dashboard. Work from the record page that owns the transaction, keep supporting documents with that record, and use preview before printing or downloading.")
    add_heading(doc, "Sign in", 2)
    add_steps(doc, ["Open the approved Procurement Sys URL.", "Enter your email and password. Complete verification or invitation acceptance if prompted.", "Confirm that the displayed organization and name are correct.", "Choose the module from the sidebar or use the dashboard shortcut."])
    add_heading(doc, "When the screen is locked", 2)
    add_text(doc, "Enter the same account password to unlock. Procurement Sys is designed to return you to the page where you were working. If the unlock returns to login, sign in again and report the page and time if it repeats.")
    add_heading(doc, "Choose your role path", 2)
    add_table(doc, ["Role", "Start with"], [
        ("Super Admin", "Dashboard, organization, users/roles, settings, backups, policies, notifications"),
        ("Procurement", "Purchase orders, suppliers, receiving follow-up, reports"),
        ("Store / workshop", "Inventory, locations, goods receiving, issues, transfers, counts"),
        ("Accountant", "Purchase orders, payment verification, petty cash, reports"),
        ("Department manager", "Department workspace, requests, allocation"),
        ("Supplier", "Supplier portal, purchase orders, products & prices, documents"),
    ], widths=[1.5, 5.4], header_fill=BLUE)
    add_callout(doc, "REMEMBER", "Do not share credentials to get around a permission limit. Ask an administrator to assign the correct role or permission.", fill="FFF7E8", accent=AMBER)


def section_structure(doc):
    add_heading(doc, "2. System structure", 1)
    add_text(doc, "Procurement Sys is one integrated platform. Each layer has a different job, but they share the same authentication, organization scope, permissions, records, and audit trail.")
    add_table(doc, ["Layer", "What it does"], [
        ("Web interface", "React, TypeScript, and responsive pages for dashboards, forms, tables, previews, and workflows."),
        ("API", "Laravel routes, controllers, validation, authorization, business rules, notifications, and synchronization."),
        ("Database", "MySQL records for organizations, users, roles, products, suppliers, orders, stock, documents, comments, and audit data."),
        ("File storage", "Uploaded signatures, supplier documents, attachments, media, and generated records. Public storage linking is required after deployment."),
        ("Background work", "Queued email and notification delivery, retries, imports, synchronization, and failure logging."),
        ("Offline client", "Protected local cache and operation queue for explicitly supported stock workflows when connectivity is unavailable."),
    ], widths=[1.55, 5.35], header_fill=BLUE_DARK)
    add_heading(doc, "Core record relationships", 2)
    add_text(doc, "A common flow is: organization -> department / location -> request -> purchase order -> supplier -> receiving -> stock movement -> payment evidence -> report. Supplier accounts connect only to their assigned supplier and permitted records. Documents and comments stay attached to the record they explain.")
    add_heading(doc, "Where to find major areas", 2)
    add_table(doc, ["Sidebar group", "Main pages"], [
        ("Overview", "Dashboard, Calendar"),
        ("Administration", "Users, Roles & Permissions, User Access"),
        ("Organization", "Companies, branches, departments, locations"),
        ("Inventory & Stores", "Products, warehouses/stores, customers"),
        ("Supplier Management", "Suppliers and supplier portal"),
        ("Procurement & Logistics", "Purchase orders, petty cash, logistics"),
        ("Media & Content", "Media Library and Offline Sync"),
        ("System & Content", "System Settings and Policies"),
    ], widths=[2.0, 4.9], header_fill=BLUE)


def section_security(doc):
    add_heading(doc, "3. Security and account safety", 1)
    add_text(doc, "Security is enforced by the backend as well as the user interface. Hiding a button is not the security control; every protected API action checks the signed-in user, organization, role, permission, and record scope.")
    add_heading(doc, "Controls used by the system", 2)
    add_bullets(doc, [
        "Bearer API sessions with expiry; the system invalidates the previous browser token when a new login is created for the account.",
        "Email verification and active-account checks before access is granted.",
        "Role-based permissions with organization-scoped roles and optional user-specific overrides.",
        "Backend authorization for viewing, creating, approving, paying, uploading, importing, synchronizing, and deleting.",
        "Input validation, upload type and size checks, rate limits on sensitive endpoints, and audit logging for important API activity.",
        "Organization and supplier scoping so users cannot use an ID from another organization to access its records.",
        "Confirmation warnings for destructive actions, including supplier and company deletion choices.",
        "Unread/read notification state, delivery status, retries, and failure logging for operational alerts.",
    ])
    add_heading(doc, "Safe daily practice", 2)
    add_bullets(doc, [
        "Use a unique password and never share it. Sign out on shared devices.",
        "Lock the screen when stepping away. Do not leave an unlocked procurement or finance page unattended.",
        "Check supplier, location, quantity, price, tax, amount, and supporting evidence before submitting or approving.",
        "Upload only business-relevant files. Use clear filenames and do not upload passwords or unrelated personal data.",
        "Use HTTPS in production and restrict database, storage, queue, and backup credentials to administrators.",
    ])
    add_callout(doc, "CRITICAL", "A super admin can delete companies or related data. Review linked records, make a backup, and confirm the exact scope before any destructive action.", fill="FCECEC", accent="B42318")


def section_navigation(doc):
    add_heading(doc, "4. Navigation and dashboard", 1)
    add_text(doc, "The sidebar shows only modules available to the current account. Search, filters, pagination, and empty states help you find records without loading the entire system at once.")
    add_heading(doc, "Dashboard and calendar", 2)
    add_bullets(doc, [
        "Dashboard values are role-aware and based on live organization data. Empty panels normally mean the underlying records have not been entered yet.",
        "Calendar shows events relevant to the account, role, organization, and procurement activity. It is not a generic holiday or leave calendar.",
        "Notifications show operational alerts such as low stock, critical stock, expiring documents, orders awaiting approval, pending deliveries, failed jobs, and synchronization problems.",
        "Use notification preferences for supported non-critical channels. Mandatory critical alerts remain controlled by administrators.",
    ])
    add_heading(doc, "Common controls", 2)
    add_table(doc, ["Control", "Use"], [
        ("Eye / View", "Open the record page or document preview, depending on the page."),
        ("Print", "Print a supported PDF from the preview header."),
        ("Download", "Save a generated PDF or uploaded file."),
        ("Filters", "Reduce the list before changing page size or searching again."),
        ("Warning modal", "Confirm the exact scope of deletion or another irreversible action."),
        ("Refresh", "Reload current data after a long-running operation or deployment."),
    ], widths=[1.55, 5.35], header_fill=BLUE)


def section_org_access(doc):
    add_heading(doc, "5. Organization and users", 1)
    add_text(doc, "Organization structure gives every transaction a clear owner, destination, and scope. Set it up before importing products or creating purchase orders.")
    add_heading(doc, "Recommended setup order", 2)
    add_steps(doc, ["Confirm company name, address, contacts, currency, and active status.", "Create branches and identify the head office.", "Create departments and assign their branch and code.", "Create warehouses and stores with unique codes.", "Create users, assign roles, and set organization, department, or supplier assignment.", "Review role access with a test account before daily operations begin."])
    add_heading(doc, "Roles and permissions", 2)
    add_text(doc, "Roles bundle permissions such as inventory.view, purchase_orders.approve, suppliers.view, supplier_catalog.manage, supplier_documents.upload, calendar.view, offline.sync, and users.manage. Super Admin has full administrative authority. Grant the smallest permission set that allows the job.")
    add_callout(doc, "ACCESS CHECK", "If a user cannot see a page or action, check active status, organization, role, permission overrides, supplier assignment, and whether the record status allows the action.", fill=BLUE_LIGHT, accent=BLUE)


def section_suppliers(doc):
    add_heading(doc, "6. Suppliers and supplier portal", 1)
    add_text(doc, "Supplier records bring legal details, contacts, commercial terms, catalog products, partnership documents, and purchase-order history together.")
    add_heading(doc, "Supplier workspace", 2)
    add_table(doc, ["Area", "Purpose"], [
        ("Profile", "Legal name, address, tax identifiers, contacts, status, and notes."),
        ("Products & prices", "Supplier products, supplier SKU, internal mapping, unit, price, lead time, and active status."),
        ("Documents", "Private contracts, certificates, insurance, agreements, and partnership files."),
        ("Purchase orders", "Orders assigned to the supplier, status, items, documents, comments, and history."),
    ], widths=[1.8, 5.1], header_fill=BLUE)
    add_heading(doc, "Supplier accounts", 2)
    add_text(doc, "A supplier-linked user sees the Supplier workspace in the sidebar: Purchase orders, Products & prices, and Documents. Access is restricted to the assigned supplier. Suppliers can update approved product and price information whenever permitted by their role.")
    add_heading(doc, "Delete supplier safely", 2)
    add_text(doc, "Use the trash icon in the supplier list. The warning flow checks linked purchase orders and offers the available safe choice: remove the supplier while preserving purchase orders and related records, or remove the supplier and its entire history when explicitly authorized. Review the scope before confirming.")


def section_procurement(doc):
    add_heading(doc, "7. Procurement and LPOs", 1)
    add_text(doc, "Purchase orders are controlled records. The available actions depend on status and permission, and the detail page is the central place for order data, attachments, comments, approvals, receiving, and payment information.")
    add_heading(doc, "Create and approve an order", 2)
    add_steps(doc, ["Select the supplier, delivery location, expected date, reason, and allocation.", "Add each product, quantity, unit cost, and tax rate. Review Qty x Unit Price = Line Total.", "Save the draft and check supplier information, items, tax, currency, totals, and supporting documents.", "Submit for approval or approve it according to your role. Do not approve your own request unless policy explicitly allows it.", "Send or issue the approved order to the supplier and follow the status from the order detail page."])
    add_heading(doc, "LPO PDF", 2)
    add_bullets(doc, [
        "The generated LPO uses the company header, preserved logo proportions, A4 layout, supplier and PO information blocks, dynamic rows, separate subtotal/VAT/grand total, and approval signatures where configured.",
        "Totals appear immediately after the last item row. Long supplier names and descriptions wrap rather than overlap.",
        "The eye icon opens the in-system preview. Print and download controls are in the preview header.",
        "The LPO detail page remains the source of truth for comments, attachments, approval state, receiving, and payment history.",
    ])
    add_heading(doc, "Receiving and payment", 2)
    add_steps(doc, ["Open the purchase-order detail page and compare the delivery against each line.", "Record received quantities and upload delivery evidence when required.", "Confirm the received quantity and stock movement.", "An authorized accountant records and verifies payment against the approved order and evidence."])
    add_callout(doc, "CONTROL", "Never use a petty cash voucher to bypass a purchase order requirement. Keep the supplier invoice, receipt, delivery evidence, and payment proof attached to the correct record.", fill="FFF7E8", accent=AMBER)


def section_petty_cash(doc):
    add_heading(doc, "8. Petty cash", 1)
    add_text(doc, "Petty Cash records small cash issues with a clear purpose, beneficiary, amount, currency, issuer, collector, and supporting evidence.")
    add_steps(doc, ["Open Petty Cash and choose Issue Petty Cash.", "Enter date, beneficiary / collector, reason, amount, currency, and notes.", "Confirm the amount is positive and the reason is specific.", "Submit and open the voucher detail page to review comments and attached documents.", "Preview, print, or download the generated voucher PDF."])
    add_heading(doc, "Voucher PDF", 2)
    add_bullets(doc, [
        "Voucher number and date are shown together in a compact metadata row.",
        "Beneficiary information is separated from the description and amount table.",
        "The total is prominent and Issued By / Collected By signature blocks are symmetrical.",
        "A visual signature can be loaded from the user profile when the account has a valid PNG signature.",
    ])


def section_inventory(doc):
    add_heading(doc, "9. Inventory and workshop", 1)
    add_text(doc, "Inventory is driven by master data and the stock ledger. A locator helps users find a part, but stock transactions are what change quantity.")
    add_heading(doc, "Master data", 2)
    add_bullets(doc, ["Create one product per SKU with a clear name, unit, category, reorder level, cost, selling price, tax rate, and active status.", "Keep product, category, unit, supplier, warehouse, and store codes consistent.", "Deactivate records instead of creating duplicates when historical transactions depend on them."])
    add_heading(doc, "Stock operations", 2)
    add_table(doc, ["Operation", "Use it for", "Check before saving"], [
        ("Receive", "Goods arrive against an order.", "Order, location, quantity, evidence."),
        ("Issue", "Goods leave a store or warehouse.", "Product, source, collector, quantity."),
        ("Transfer", "Goods move between authorized locations.", "Dispatch location, receiving location, quantity."),
        ("Adjust", "Approved discrepancy or opening balance.", "Reason, product, location, quantity."),
        ("Count", "Physical stock comparison.", "Count date, location, variance follow-up."),
    ], widths=[1.0, 3.2, 2.7], header_fill=BLUE_DARK)
    add_heading(doc, "Workshop warehouse", 2)
    add_text(doc, "Workshop Issues records issued goods and can generate a document preview. The Spare Part Locator helps find stock by product or SKU. It does not replace the stock ledger, issue, transfer, or adjustment workflow.")


def section_documents(doc):
    add_heading(doc, "10. Documents and previews", 1)
    add_text(doc, "Documents can be viewed directly inside Procurement Sys. Use the preview for reading and the record detail page for workflow actions, comments, metadata, and permissions.")
    add_heading(doc, "Where previews are available", 2)
    add_bullets(doc, ["Purchase orders and supplier portal purchase orders.", "Petty cash vouchers.", "Goods received and issued workshop warehouse documents.", "Media Library files and supplier partnership documents.", "Record attachments such as receipts, invoices, proof of payment, and delivery evidence."])
    add_heading(doc, "Preview controls", 2)
    add_table(doc, ["Action", "What happens"], [
        ("Eye icon", "Opens a file in view mode inside the system."),
        ("Print", "Prints a supported document from the viewer header."),
        ("Download", "Downloads the original or generated file."),
        ("Close", "Returns to the underlying page without losing the record context."),
    ], widths=[1.4, 5.5], header_fill=BLUE)
    add_callout(doc, "SUPPORTED FILES", "PDFs, images, and browser-readable text files can be previewed. Office or archive files may offer download only when the browser cannot render them safely.", fill=BLUE_LIGHT, accent=BLUE)
    add_heading(doc, "Signature image", 2)
    add_text(doc, "Users can upload a transparent-background PNG signature from Profile. The signature is stored with the user profile and inserted visually into supported generated PDFs at the configured signature field. This is a visual signature, not a certificate-based cryptographic signature.")


def section_offline(doc):
    add_heading(doc, "11. Offline mode and synchronization", 1)
    add_text(doc, "Offline mode is for approved operations in unreliable-connectivity environments. It is not an unrestricted copy of the system. The local cache contains only the minimum data allowed for the user and configured workflow.")
    add_heading(doc, "Before going offline", 2)
    add_bullets(doc, ["Sign in while connected and confirm the account, organization, role, permissions, locations, products, and prices are current.", "Open Offline Sync and confirm the last successful synchronization.", "Do not use a shared device or leave sensitive records cached after the session ends."])
    add_heading(doc, "While offline", 2)
    add_text(doc, "Supported operations may include stock counts, receiving, issuing, transfers, or other configured actions. Each local transaction receives a unique client operation ID and remains in the local queue until the server confirms it.")
    add_heading(doc, "Queue states", 2)
    add_table(doc, ["State", "Meaning"], [
        ("Pending", "Created locally and waiting for a connection."),
        ("Syncing", "Being sent to the server."),
        ("Synchronized", "Accepted by the server; safe to treat as posted."),
        ("Failed", "Could not be accepted; read the message and retry after fixing the cause."),
        ("Conflict", "Server data changed or the operation is no longer valid; resolve it instead of overwriting stock."),
    ], widths=[1.45, 5.45], header_fill=BLUE_DARK)
    add_heading(doc, "When connection returns", 2)
    add_steps(doc, ["Open Offline Sync and start synchronization or wait for the configured reconnect attempt.", "Review successful, failed, and conflict operations.", "Retry transient failures after checking permissions, product status, quantity, and connectivity.", "For a conflict, follow the message and ask an administrator when the business decision cannot be made locally."])
    add_callout(doc, "IMPORTANT", "Retrying the same client operation ID is idempotent: it must not create duplicate receipts, issues, transfers, or reservations. Do not manually recreate a failed operation until its queue state is understood.", fill="FFF7E8", accent=AMBER)


def section_admin(doc):
    add_heading(doc, "12. Policies and administration", 1)
    add_heading(doc, "Policies page", 2)
    add_text(doc, "The Policies page displays four organization policies inside the system: Terms of Service, Privacy Policy, Supplier Policy, and Procurement Ordering Policy. All authenticated users can read published policies. Super Admin can edit title, summary, content, and publication status.")
    add_steps(doc, ["Open Policies from System & Content.", "Select a policy from the library.", "Super Admin chooses Edit policy, updates the content, and saves.", "Publish the policy when it is ready for users to read."])
    add_heading(doc, "System settings", 2)
    add_text(doc, "Authorized administrators can manage organization settings, email configuration, cache clearing, backups, notification rules, and other system controls. Keep production settings aligned with the hosting environment and never expose database credentials in screenshots or manuals.")
    add_heading(doc, "Backups and recovery", 2)
    add_bullets(doc, ["Confirm the database host, port, database name, username, and password in the hosting environment before creating a backup.", "Store backups outside the public web root and test that a backup can be restored.", "Backups may fail when MySQL is stopped, the host is wrong, the user lacks export permission, or a server tool is unavailable.", "Use the backup before broad company, supplier, or data deletion."])
    add_heading(doc, "Notifications administration", 2)
    add_text(doc, "Administrators can enable or disable operational event rules, choose severity, configure recipients and channels, and enforce mandatory critical alerts. Users can manage supported non-critical preferences from Notifications. Email delivery is processed asynchronously so a provider delay does not block an inventory transaction.")


def section_troubleshooting(doc):
    add_heading(doc, "13. Troubleshooting", 1)
    add_table(doc, ["Message / symptom", "First response"], [
        ("Page will not display", "Refresh once, confirm the API is running, check permissions, and retry the direct route. Report the time and URL if it persists."),
        ("Unlock returns to login", "Sign in again. If it repeats, report the page, time, browser, and account role."),
        ("Database export failed", "Check that MySQL is running and that the configured database host is resolvable. Verify credentials and export permissions."),
        ("Unknown column in products", "The database schema is behind the application. Run the pending Laravel migrations and review the migration output."),
        ("Maximum call stack size exceeded", "Stop repeated sync clicks, refresh the page, and retry once. If it repeats, capture the operation and browser console error."),
        ("Document is upside down", "Refresh the application and reopen the file. If it is a generated PDF, report the document number so the source PDF can be checked."),
        ("Preview is narrow or black", "Hard-refresh, reset browser zoom, and reopen the viewer. Use print/download if the browser viewer cannot render the file."),
        ("Signature rejected", "Select a real PNG signature file within the allowed size, hard-refresh, and upload again. Renaming a JPG to .png does not convert it."),
        ("Cannot delete supplier", "Review linked PO history and choose whether to preserve records or remove the complete history, if authorized."),
        ("Offline conflict", "Do not overwrite it. Read the conflict message, review server quantity/status, and ask the responsible administrator when needed."),
    ], widths=[2.1, 4.8], header_fill=BLUE_DARK)
    add_callout(doc, "REPORT WITH CONTEXT", "When contacting support, include your account role, organization, page URL, record number, exact message, approximate time, and whether you were online or offline. Do not include passwords or secret tokens.", fill=BLUE_LIGHT, accent=BLUE)
    add_heading(doc, "Glossary", 2)
    add_table(doc, ["Term", "Meaning"], [
        ("LPO", "Local Purchase Order generated from an approved purchase order record."),
        ("Goods receipt", "Recorded quantity physically received against a purchase order."),
        ("Supplier account", "Supplier-linked user with access limited to the assigned supplier workspace."),
        ("Media asset", "Organization file stored in Media Library."),
        ("Operation ID", "Unique client identifier used to synchronize an offline transaction safely."),
        ("Permission override", "User-specific allow or deny that changes the effective role permission."),
    ], widths=[1.65, 5.25], header_fill=BLUE)
    add_text(doc, "End of manual", before=16, after=0, color=MUTED, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)
    add_contents(doc)
    section_quick_start(doc)
    doc.add_page_break()
    section_structure(doc)
    doc.add_page_break()
    section_security(doc)
    doc.add_page_break()
    section_navigation(doc)
    doc.add_page_break()
    section_org_access(doc)
    doc.add_page_break()
    section_suppliers(doc)
    doc.add_page_break()
    section_procurement(doc)
    doc.add_page_break()
    section_petty_cash(doc)
    doc.add_page_break()
    section_inventory(doc)
    doc.add_page_break()
    section_documents(doc)
    doc.add_page_break()
    section_offline(doc)
    doc.add_page_break()
    section_admin(doc)
    doc.add_page_break()
    section_troubleshooting(doc)
    doc.core_properties.title = "Procurement Sys User Manual"
    doc.core_properties.subject = "Concise user and administrator guide"
    doc.core_properties.author = "Procurement Sys"
    doc.core_properties.comments = "Updated user manual"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
